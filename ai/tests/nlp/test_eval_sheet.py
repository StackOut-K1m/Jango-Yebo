"""라벨링 시트 왕복. 여기서 유실이 나면 사람이 몇 시간 채운 라벨이 날아간다."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from evals import labels, sheet
from evals.labels import LabelRecord

RAW = {
    "transactionId": 1,
    "description": "김영자",
    "direction": "WITHDRAWAL",
    "amount": 1200000,
    "occurredAt": "2026-08-25T01:00:00Z",
    "cycleHint": "매월 25일 반복 (최근 4회)",
    "industryType": "FOOD_CAFE",
    "source": "synthetic-20260908",
}

DEPOSIT = {
    **RAW,
    "transactionId": 2,
    "description": "우아한형제들",
    "direction": "DEPOSIT",
    "amount": 748000,
    "cycleHint": None,
}


def _record(**overrides) -> LabelRecord:
    return LabelRecord.model_validate({**RAW, **overrides})


@pytest.fixture
def pool() -> list[LabelRecord]:
    return [_record(), LabelRecord.model_validate(DEPOSIT)]


def _set(path: Path, row: int, **cells) -> None:
    """시트의 라벨 칸을 사람이 채운 것처럼 바꾼다.

    `.value` 에 직접 넣는다 — `ws.cell(..., value=None)` 은 내부에서
    `if value is not None` 으로 걸러서 셀을 비우지 못한다.
    """
    columns = {
        "category": sheet.CATEGORY_COL,
        "merchant": sheet.MERCHANT_COL,
        "note": sheet.NOTE_COL,
    }
    wb = load_workbook(path)
    ws = wb[sheet.SHEET_NAME]
    for key, value in cells.items():
        ws.cell(row=row, column=columns[key]).value = value
    wb.save(path)


# ---------------------------------------------------------------- 왕복


def test_roundtrip_carries_labels_back(tmp_path: Path, pool: list[LabelRecord]):
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)

    _set(xlsx, 2, category="RENT", merchant="김영자(임대인 추정)", note="매월 25일 동일 금액")
    updated, _ = sheet.do_import(pool, xlsx)

    first = next(r for r in updated if r.transaction_id == 1)
    assert first.category == "RENT"
    assert first.merchant_name == "김영자(임대인 추정)"
    assert first.note == "매월 25일 동일 금액"


def test_unfilled_rows_stay_unlabeled(tmp_path: Path, pool: list[LabelRecord]):
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)

    updated, summary = sheet.do_import(pool, xlsx)
    assert all(not r.labeled for r in updated)
    assert "새로 채운 것 0건" in summary[0]


def test_emptying_a_cell_clears_the_label(tmp_path: Path, pool: list[LabelRecord]):
    """실수로 붙인 라벨을 시트에서 지우면 jsonl 에서도 지워져야 한다."""
    labeled = [_record(category="RENT"), pool[1]]
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(labeled, xlsx)

    _set(xlsx, 2, category=None)
    updated, summary = sheet.do_import(labeled, xlsx)

    assert next(r for r in updated if r.transaction_id == 1).category is None
    assert any("지움" in line for line in summary)


def test_export_preserves_existing_labels(tmp_path: Path, pool: list[LabelRecord]):
    """중간에 끊고 다시 내보낼 때 이미 채운 것이 시트에 남아 있어야 한다."""
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export([_record(category="RENT", merchantName="김영자"), pool[1]], xlsx)

    ws = load_workbook(xlsx)[sheet.SHEET_NAME]
    assert ws.cell(row=2, column=sheet.CATEGORY_COL).value == "RENT"
    assert ws.cell(row=2, column=sheet.MERCHANT_COL).value == "김영자"


# ---------------------------------------------------------------- 정본 보호


def test_import_ignores_edits_to_context_columns(tmp_path: Path, pool: list[LabelRecord]):
    """Excel 이 날짜·금액 서식을 제 맘대로 바꿔도 정본이 다치지 않아야 한다.

    맥락 열은 labels.jsonl 에서 다시 읽는다. 시트는 라벨 3칸만 공급한다.
    """
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)

    wb = load_workbook(xlsx)
    ws = wb[sheet.SHEET_NAME]
    ws.cell(row=2, column=3, value="엉뚱한 적요")  # 적요
    ws.cell(row=2, column=5, value=99)  # 금액
    ws.cell(row=2, column=4, value="DEPOSIT")  # 방향
    wb.save(xlsx)

    updated, _ = sheet.do_import(pool, xlsx)
    first = next(r for r in updated if r.transaction_id == 1)
    assert first.description == "김영자"
    assert first.amount == 1200000
    assert first.direction == "WITHDRAWAL"


def test_direction_mismatch_is_rejected(tmp_path: Path, pool: list[LabelRecord]):
    """드롭다운을 무시하고 손으로 타이핑하면 여기서 걸려야 한다."""
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)
    _set(xlsx, 2, category="PLATFORM_SETTLEMENT")  # 출금 거래에 입금 카테고리

    with pytest.raises(ValueError, match="transactionId=1"):
        sheet.do_import(pool, xlsx)


def test_unknown_transaction_id_is_rejected(tmp_path: Path, pool: list[LabelRecord]):
    """시트에 행을 새로 추가하면 정답셋과 대응이 깨진다."""
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)

    wb = load_workbook(xlsx)
    ws = wb[sheet.SHEET_NAME]
    ws.cell(row=99, column=1, value=12345)
    wb.save(xlsx)

    with pytest.raises(ValueError, match="없는 transactionId"):
        sheet.do_import(pool, xlsx)


def test_wrong_sheet_name_is_rejected(tmp_path: Path, pool: list[LabelRecord]):
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)

    wb = load_workbook(xlsx)
    wb[sheet.SHEET_NAME].title = "Sheet1"
    wb.save(xlsx)

    with pytest.raises(ValueError, match="시트가 없습니다"):
        sheet.do_import(pool, xlsx)


def test_lowercase_typed_category_is_accepted(tmp_path: Path, pool: list[LabelRecord]):
    """손으로 치면 소문자로 칠 수 있다. 그것 때문에 실패시킬 이유는 없다."""
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)
    _set(xlsx, 2, category=" rent ")

    updated, _ = sheet.do_import(pool, xlsx)
    assert next(r for r in updated if r.transaction_id == 1).category == "RENT"


# ---------------------------------------------------------------- 드롭다운


def test_dropdown_offers_only_matching_direction(tmp_path: Path, pool: list[LabelRecord]):
    """출금 거래에 입금 카테고리를 고를 수 있으면 안 된다.

    labels.load() 가 나중에 잡아 주지만, 300건 채운 뒤에 알게 되면 늦다.
    """
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)

    lists = load_workbook(xlsx)[sheet.LIST_SHEET]
    deposits = [c.value for c in lists["A"] if c.value][1:]
    withdrawals = [c.value for c in lists["B"] if c.value][1:]

    assert len(deposits) == 7 and len(withdrawals) == 16
    assert set(deposits).isdisjoint(withdrawals)
    assert "RENT" in withdrawals and "RENT" not in deposits


def test_list_sheet_is_hidden(tmp_path: Path, pool: list[LabelRecord]):
    """드롭다운 원본이 보이면 라벨러가 실수로 고친다."""
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)
    assert load_workbook(xlsx)[sheet.LIST_SHEET].sheet_state == "hidden"


def test_context_columns_are_frozen(tmp_path: Path, pool: list[LabelRecord]):
    """오른쪽으로 스크롤해도 적요·금액이 보여야 판단할 수 있다."""
    xlsx = tmp_path / "labels.xlsx"
    sheet.do_export(pool, xlsx)
    assert load_workbook(xlsx)[sheet.SHEET_NAME].freeze_panes == "I2"


@pytest.mark.parametrize(
    ("rows", "expected"),
    [
        ([2], "I2"),
        ([2, 3, 4], "I2:I4"),
        ([2, 3, 5], "I2:I3 I5"),
        ([5, 2, 3], "I2:I3 I5"),
        ([2, 4, 6], "I2 I4 I6"),
    ],
)
def test_collapse_builds_contiguous_ranges(rows: list[int], expected: str):
    """셀 300개를 낱개로 나열하면 파일이 커지고 Excel 이 느려진다."""
    assert sheet._collapse(rows, sheet.CATEGORY_COL) == expected


# ---------------------------------------------------------------- CLI


def test_export_refuses_to_overwrite(tmp_path: Path, pool: list[LabelRecord], capsys):
    """시트에 채운 라벨을 덮어쓰는 사고를 막는다."""
    dataset = tmp_path / "p9-99"
    labels.dump(pool, dataset / "labels.jsonl")

    monkey = sheet.EVALS_ROOT
    try:
        sheet.EVALS_ROOT = tmp_path
        assert sheet.main(["export", "p9-99"]) == 0
        assert sheet.main(["export", "p9-99"]) == 1
        assert "이미 있습니다" in capsys.readouterr().out
        assert sheet.main(["export", "p9-99", "--force"]) == 0
    finally:
        sheet.EVALS_ROOT = monkey


def test_import_leaves_jsonl_untouched_on_violation(tmp_path: Path, pool: list[LabelRecord]):
    """가져오다 실패했을 때 정본이 반쯤 덮어써지면 안 된다."""
    dataset = tmp_path / "p9-99"
    jsonl = dataset / "labels.jsonl"
    labels.dump(pool, jsonl)
    before = jsonl.read_bytes()

    monkey = sheet.EVALS_ROOT
    try:
        sheet.EVALS_ROOT = tmp_path
        assert sheet.main(["export", "p9-99"]) == 0
        _set(dataset / "labels.xlsx", 2, category="PLATFORM_SETTLEMENT")
        assert sheet.main(["import", "p9-99"]) == 1
    finally:
        sheet.EVALS_ROOT = monkey

    assert jsonl.read_bytes() == before
