"""라벨링용 xlsx 를 내보내고 다시 읽어들인다. 300건을 손으로 채우는 작업이라
드롭다운으로 오타와 방향 불일치를 원천 차단하는 게 목적이다.

    uv run --frozen python -m evals.sheet export p1-01   # jsonl → xlsx
    (Excel 에서 category·거래처명·메모 를 채운다)
    uv run --frozen python -m evals.sheet import p1-01   # xlsx → jsonl

**정본은 labels.jsonl 이다.** xlsx 는 입력 도구일 뿐이고, git 에도 jsonl 만 올라간다
(바이너리는 diff 가 안 읽혀서 라벨 수정 이력을 리뷰할 수 없다).

xlsx 를 쓰는 것은 CSV 와 다르다 — xlsx 내부는 UTF-8 XML 이라 Excel 이 저장해도
한글이 깨지지 않는다. cp949 재인코딩 사고는 .csv 확장자에서만 난다.

`import` 는 xlsx 에서 **라벨 3칸만** 가져온다. 거래일시·금액 같은 맥락 열은 jsonl 에서
그대로 다시 읽는다 — Excel 이 날짜를 제 형식으로 바꿔 놓아도 정본이 다치지 않게.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import get_args

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.nlp.schemas import CATEGORY_DIRECTION, TransactionCategory
from evals import labels
from evals.labels import LabelRecord

EVALS_ROOT = Path(__file__).resolve().parent

SHEET_NAME = "라벨링"
LIST_SHEET = "_lists"  # 드롭다운 원본. 시트명을 ASCII 로 둔 것은 수식 참조 호환 때문이다

# 열 순서. 앞 8개는 읽기 전용 맥락이고, 뒤 3개를 사람이 채운다.
CONTEXT_HEADERS = [
    ("transactionId", 14),
    ("거래일시", 18),
    ("적요", 22),
    ("방향", 12),
    ("금액", 12),
    ("주기힌트", 26),
    ("업종", 14),
    ("출처", 20),
]
LABEL_HEADERS = [("category", 22), ("거래처명", 24), ("메모", 40)]

FIRST_LABEL_COL = len(CONTEXT_HEADERS) + 1  # I
CATEGORY_COL = FIRST_LABEL_COL
MERCHANT_COL = FIRST_LABEL_COL + 1
NOTE_COL = FIRST_LABEL_COL + 2

_HEADER_FILL = PatternFill("solid", fgColor="FF2F4858")
_CONTEXT_FILL = PatternFill("solid", fgColor="FFF2F2F2")
_INPUT_FILL = PatternFill("solid", fgColor="FFFFF9E6")


def _categories(direction: str) -> list[str]:
    """방향에 맞는 카테고리만. 입금 7종 / 출금 16종으로 갈린다."""
    return [c for c in get_args(TransactionCategory) if CATEGORY_DIRECTION[c] == direction]


def _collapse(rows: list[int], column: int) -> str:
    """연속한 행을 구간으로 묶어 sqref 문자열을 만든다 ("I2:I4 I7").

    셀을 300개 낱개로 나열해도 Excel 이 받기는 하지만, 파일이 커지고 열 때 느려진다.
    """
    letter = get_column_letter(column)
    parts: list[str] = []
    start = prev = None
    for row in sorted(rows):
        if start is None:
            start = prev = row
        elif row == prev + 1:
            prev = row
        else:
            parts.append(f"{letter}{start}" if start == prev else f"{letter}{start}:{letter}{prev}")
            start = prev = row
    if start is not None:
        parts.append(f"{letter}{start}" if start == prev else f"{letter}{start}:{letter}{prev}")
    return " ".join(parts)


# ------------------------------------------------------------------ 내보내기


def do_export(records: list[LabelRecord], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    for col, (name, width) in enumerate(CONTEXT_HEADERS + LABEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    by_direction: dict[str, list[int]] = {"DEPOSIT": [], "WITHDRAWAL": []}

    for offset, r in enumerate(sorted(records, key=lambda r: r.transaction_id)):
        row = offset + 2
        values = [
            r.transaction_id,
            r.occurred_at.replace(tzinfo=None),
            r.description,
            r.direction,
            r.amount,
            r.cycle_hint or "",
            r.industry_type,
            r.source,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.fill = _CONTEXT_FILL
        ws.cell(row=row, column=2).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=row, column=5).number_format = "#,##0"

        ws.cell(row=row, column=CATEGORY_COL, value=r.category or "").fill = _INPUT_FILL
        ws.cell(row=row, column=MERCHANT_COL, value=r.merchant_name or "").fill = _INPUT_FILL
        note = ws.cell(row=row, column=NOTE_COL, value=r.note or "")
        note.fill = _INPUT_FILL
        note.alignment = Alignment(wrap_text=True, vertical="top")

        by_direction[r.direction].append(row)

    _write_dropdowns(wb, ws, by_direction)

    # 맥락 열을 고정해 오른쪽으로 스크롤해도 적요·금액이 보이게 한다.
    ws.freeze_panes = f"{get_column_letter(FIRST_LABEL_COL)}2"
    # 방향으로 걸러 같은 종류를 몰아서 라벨링하면 훨씬 빠르다.
    ws.auto_filter.ref = f"A1:{get_column_letter(NOTE_COL)}{len(records) + 1}"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_dropdowns(wb: Workbook, ws, by_direction: dict[str, list[int]]) -> None:
    """방향별로 다른 드롭다운을 건다.

    한 목록에 23종을 몰아넣으면 출금 거래에 입금 카테고리를 고를 수 있게 된다.
    그건 labels.load() 가 잡아 주지만, 애초에 고를 수 없게 하는 편이 낫다.
    """
    lists = wb.create_sheet(LIST_SHEET)
    lists["A1"] = "DEPOSIT"
    lists["B1"] = "WITHDRAWAL"

    for col, direction in ((1, "DEPOSIT"), (2, "WITHDRAWAL")):
        options = _categories(direction)
        for i, category in enumerate(options, start=2):
            lists.cell(row=i, column=col, value=category)

        rows = by_direction[direction]
        if not rows:
            continue

        letter = get_column_letter(col)
        dv = DataValidation(
            type="list",
            formula1=f"{LIST_SHEET}!${letter}$2:${letter}${len(options) + 1}",
            allow_blank=True,
            errorStyle="stop",
            errorTitle="목록에 없는 값",
            error=f"{direction} 거래에는 {len(options)}종만 쓸 수 있습니다. 드롭다운에서 고르세요.",
            promptTitle=f"{direction} 카테고리",
            prompt="적요·금액·주기힌트만 보고 판단하세요. 애매하면 메모에 이유를 남기세요.",
            showInputMessage=True,
            showErrorMessage=True,
        )
        ws.add_data_validation(dv)
        dv.sqref = _collapse(rows, CATEGORY_COL)

    lists.sheet_state = "hidden"


# ------------------------------------------------------------------ 읽어들이기


def _cell_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def do_import(records: list[LabelRecord], path: Path) -> tuple[list[LabelRecord], list[str]]:
    """xlsx 의 라벨 3칸을 jsonl 레코드에 얹는다. 맥락 열은 보지 않는다.

    반환값은 (갱신된 레코드, 사람에게 보일 변경 요약).
    규격을 어기면 ValueError — 그때는 아무것도 쓰지 않는다.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"'{SHEET_NAME}' 시트가 없습니다. 시트 이름을 바꾸지 마세요.")
    ws = wb[SHEET_NAME]

    by_id = {r.transaction_id: r for r in records}
    incoming: dict[int, tuple[str | None, str | None, str | None]] = {}
    unknown: list[int] = []

    for row in ws.iter_rows(min_row=2, max_col=NOTE_COL, values_only=True):
        if row[0] is None:
            continue
        try:
            transaction_id = int(row[0])
        except (TypeError, ValueError):
            raise ValueError(f"transactionId 가 숫자가 아닙니다: {row[0]!r}") from None

        if transaction_id not in by_id:
            unknown.append(transaction_id)
            continue

        category = _cell_text(row[CATEGORY_COL - 1])
        incoming[transaction_id] = (
            category.upper() if category else None,
            _cell_text(row[MERCHANT_COL - 1]),
            _cell_text(row[NOTE_COL - 1]),
        )

    if unknown:
        raise ValueError(
            f"labels.jsonl 에 없는 transactionId {len(unknown)}건: {sorted(unknown)[:10]}\n"
            "시트에 행을 새로 추가했거나 transactionId 를 고쳤습니다. 되돌리세요."
        )

    updated: list[LabelRecord] = []
    changes: list[str] = []
    added = changed = cleared = 0

    for record in records:
        category, merchant, note = incoming.get(
            record.transaction_id, (record.category, record.merchant_name, record.note)
        )
        # model_copy 로 얹으면 검증이 다시 돌지 않는다. 사람이 채운 값이라 방향과
        # 카테고리가 안 맞을 수 있으므로, 반드시 재검증을 거쳐 만든다.
        payload = record.model_dump(by_alias=True, mode="json")
        payload.update({"category": category, "merchantName": merchant, "note": note})
        try:
            rebuilt = LabelRecord.model_validate(payload)
        except ValueError as exc:
            raise ValueError(f"transactionId={record.transaction_id} 규격 위반\n{exc}") from exc

        if record.category != rebuilt.category:
            if record.category is None:
                added += 1
            elif rebuilt.category is None:
                cleared += 1
                changes.append(f"  지움  {record.transaction_id} {record.category} → (빈칸)")
            else:
                changed += 1
                changes.append(
                    f"  변경  {record.transaction_id} {record.category} → {rebuilt.category}"
                )
        updated.append(rebuilt)

    summary = [f"새로 채운 것 {added}건", f"바꾼 것 {changed}건", f"지운 것 {cleared}건"]
    return updated, summary + changes


# ------------------------------------------------------------------ CLI


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="라벨링 xlsx 내보내기·읽어들이기")
    sub = parser.add_subparsers(dest="command", required=True)

    for name, help_text in (
        ("export", "labels.jsonl → labels.xlsx (드롭다운 포함)"),
        ("import", "labels.xlsx → labels.jsonl (라벨 3칸만 가져온다)"),
    ):
        p = sub.add_parser(name, help=help_text)
        p.add_argument("dataset", help="세트 이름 (예: p1-01)")
        p.add_argument("--file", default="labels.jsonl", help="라벨 파일명")
        p.add_argument("--xlsx", default="labels.xlsx", help="시트 파일명")
        if name == "export":
            p.add_argument("--force", action="store_true", help="기존 xlsx 를 덮어쓴다")

    args = parser.parse_args(argv)
    root = EVALS_ROOT / args.dataset
    jsonl_path = root / args.file
    xlsx_path = root / args.xlsx

    if not jsonl_path.exists():
        print(f"라벨 파일이 없습니다: {jsonl_path}")
        print("표본을 먼저 뽑으세요 — python -m evals.sample --help")
        return 1

    try:
        records = labels.load(jsonl_path)
    except ValueError as exc:
        print(f"labels.jsonl 규격 위반\n{exc}")
        return 1

    if args.command == "export":
        return _run_export(records, xlsx_path, force=args.force)
    return _run_import(records, jsonl_path, xlsx_path)


def _run_export(records: list[LabelRecord], xlsx_path: Path, *, force: bool) -> int:
    if xlsx_path.exists() and not force:
        print(f"이미 있습니다: {xlsx_path}")
        print("덮어쓰면 시트에 채운 라벨이 사라집니다.")
        print("시트의 라벨을 먼저 가져오세요 — python -m evals.sheet import <세트>")
        print("정말 새로 만들려면 --force")
        return 1

    do_export(records, xlsx_path)
    labeled = sum(1 for r in records if r.labeled)
    print(f"{len(records)}건을 {xlsx_path} 에 썼습니다 (이미 라벨된 것 {labeled}건)")
    print("category 열은 방향에 맞는 카테고리만 나오는 드롭다운입니다")
    print("\n채운 뒤:")
    print("  uv run --frozen python -m evals.sheet import <세트>")
    return 0


def _run_import(records: list[LabelRecord], jsonl_path: Path, xlsx_path: Path) -> int:
    if not xlsx_path.exists():
        print(f"시트가 없습니다: {xlsx_path}")
        print("먼저 내보내세요 — python -m evals.sheet export <세트>")
        return 1

    try:
        updated, summary = do_import(records, xlsx_path)
    except ValueError as exc:
        print(f"가져오지 못했습니다 — labels.jsonl 은 그대로입니다\n\n{exc}")
        return 1

    labels.dump(updated, jsonl_path)
    labeled = sum(1 for r in updated if r.labeled)

    print(f"{xlsx_path.name} → {jsonl_path.name}")
    for line in summary:
        print(line if line.startswith("  ") else f"  {line}")
    print(f"\n라벨링 {labeled} / {len(updated)}건 ({labeled / len(updated):.0%})")
    print("\n분포 점검:")
    print("  uv run --frozen python -m evals.check <세트>")
    return 0


if __name__ == "__main__":
    sys.exit(main())
